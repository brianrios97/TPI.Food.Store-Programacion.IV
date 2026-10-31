import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Establece color de fondo de una celda."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de una celda."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tc_pr.append(tcMar)

def create_guion_docx(output_path: str):
    doc = Document()

    # Configuración de márgenes normales (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilo base
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x29, 0x25, 0x24) # stone-800

    # --- PORTADA / ENCABEZADO ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("GUION DE GRABACIÓN - VIDEO DEMOSTRATIVO")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0xD9, 0x77, 0x06) # amber-600

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Food Store - Sistema de Pedidos y Delivery Gastronómico")
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x1C, 0x19, 0x17)

    # Tabla de Metadatos
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    metadata = [
        ("Materia:", "Programación IV - Tecnicatura Universitaria en Programación de Sistemas (UTN)"),
        ("Autor / Alumno:", "Brian Emanuel Rios"),
        ("Repositorio GitHub:", "https://github.com/brianrios97/TPI.Food.Store-Programacion.IV"),
        ("Duración Estimada:", "7 a 10 minutos (Ritmo pausado y claro)")
    ]

    for idx, (label, val) in enumerate(metadata):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.5)

        set_cell_background(cell_lbl, "F5F5F4")
        set_cell_background(cell_val, "FAFAF9")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        r_l = p_lbl.add_run(label)
        r_l.bold = True
        r_l.font.size = Pt(9.5)

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        r_v = p_val.add_run(val)
        r_v.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECCIÓN: PREPARACIÓN PREVIA ---
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Preparación Previa (Antes de Iniciar la Grabación)")
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

    prep_items = [
        ("Terminal 1 (Backend):", "Ubicada en la carpeta backend con el entorno virtual activado (.venv\\Scripts\\activate)."),
        ("Terminal 2 (Frontend):", "Ubicada en la carpeta frontend lista para ejecutar pnpm build y pnpm dev."),
        ("Pestañas del Navegador:", "Pestaña 1 en http://localhost:5173 (Food Store), Pestaña 2 en http://localhost:8000/docs (Swagger), Pestaña 3 con el repositorio de GitHub."),
        ("Editor de Código:", "VS Code abierto en la raíz con los archivos unit_of_work.py, service.py y authStore.ts a mano.")
    ]

    for label, desc in prep_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(label + " ")
        r1.bold = True
        p.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Función auxiliar para agregar bloques del guion
    def add_bloque(titulo_bloque, tiempo, que_mostrar, que_decir):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(4)
        r_h = h.add_run(f"{titulo_bloque} ({tiempo})")
        r_h.font.size = Pt(13)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)

        # Caja: Qué mostrar en pantalla
        table_show = doc.add_table(rows=1, cols=1)
        table_show.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_show = table_show.rows[0].cells[0]
        cell_show.width = Inches(6.5)
        set_cell_background(cell_show, "EFF6FF") # blue-50
        set_cell_margins(cell_show, top=100, bottom=100, left=140, right=140)

        p_s = cell_show.paragraphs[0]
        p_s.paragraph_format.space_after = Pt(4)
        r_st = p_s.add_run("🖥️ QUÉ MOSTRAR EN PANTALLA:")
        r_st.bold = True
        r_st.font.size = Pt(10)
        r_st.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

        for item in que_mostrar:
            pi = cell_show.add_paragraph(style='List Bullet')
            pi.paragraph_format.space_after = Pt(2)
            ri = pi.add_run(item)
            ri.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Caja: Qué decir
        table_say = doc.add_table(rows=1, cols=1)
        table_say.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_say = table_say.rows[0].cells[0]
        cell_say.width = Inches(6.5)
        set_cell_background(cell_say, "FEF3C7") # amber-100 / warm
        set_cell_margins(cell_say, top=120, bottom=120, left=150, right=150)

        p_say = cell_say.paragraphs[0]
        p_say.paragraph_format.space_after = Pt(4)
        r_sayt = p_say.add_run("🎙️ QUÉ DECIR (GUION HABLADO):")
        r_sayt.bold = True
        r_sayt.font.size = Pt(10)
        r_sayt.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

        p_text = cell_say.add_paragraph()
        p_text.paragraph_format.space_after = Pt(0)
        p_text.paragraph_format.line_spacing = 1.15
        r_body = p_text.add_run(f'"{que_decir}"')
        r_body.font.size = Pt(10.5)
        r_body.font.italic = True

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- BLOQUES ---
    add_bloque(
        "Bloque 1: Presentación y Arquitectura General",
        "0:00 - 1:30",
        [
            "Cámara web encendida o portada del navegador con el README.md.",
            "Mostrar la estructura en capas en VS Code: app/core, app/modules (auth, categorias, pedidos, productos, usuarios).",
            "Mostrar las carpetas de frontend: src/features, src/routes, src/store, src/shared."
        ],
        "Buenas tardes a los profesores de la cátedra de Programación IV. Mi nombre es Brian Emanuel Rios y voy a presentar mi Trabajo Práctico Integrador: Food Store, una plataforma web para pedidos y delivery gastronómico.\n\n"
        "Para este proyecto tomé como punto de partida el trabajo que realicé en Programación III, evolucionando su arquitectura hacia un estándar profesional desacoplado.\n\n"
        "En el backend utilizamos Python 3.11 con FastAPI y SQLModel sobre SQLAlchemy 2. Implementamos una arquitectura limpia por capas: Router, Service, Unit of Work con context manager, repositorios genéricos y modelos de dominio. La autenticación se maneja con JWT firmado en cookies HttpOnly para evitar ataques XSS.\n\n"
        "En el frontend construimos una Single Page Application con React 19 y TypeScript estricto, sin ningún tipo any, empaquetado con Vite 7 y Tailwind CSS 4. El estado del servidor lo gestionamos con TanStack Query 5, el estado del cliente con Zustand 5 persistente, y los formularios y tablas con TanStack Form 1 y TanStack Table 8."
    )

    add_bloque(
        "Bloque 2: Puesta en Marcha, Seed y Documentación Swagger",
        "1:30 - 2:45",
        [
            "En la Terminal 1 (Backend), ejecutar: python -m app.db.seed.",
            "Mostrar el mensaje de creación de tablas, usuarios (admin/cliente) y catálogo sin duplicados.",
            "Iniciar FastAPI: uvicorn app.main:app --reload --port 8000.",
            "En el navegador abrir http://localhost:8000/docs y mostrar los endpoints organizados.",
            "En la Terminal 2 (Frontend), ejecutar pnpm build para demostrar 0 errores de TypeScript, y luego pnpm dev."
        ],
        "Comenzamos mostrando la inicialización del sistema. En la consola del backend ejecutamos 'python -m app.db.seed'. Este script es completamente idempotente: verifica si los usuarios base o los productos ya existen y no genera duplicados, dejando la base lista para operar.\n\n"
        "Iniciamos el servidor con Uvicorn en el puerto 8000 y vemos en el navegador la documentación interactiva Swagger UI. Aquí se observan los endpoints modulares de autenticación, catálogo, pedidos y administración.\n\n"
        "En la terminal del frontend ejecutamos 'pnpm build' que corre 'tsc -b' y empaqueta con Vite. Como pueden comprobar, la compilación finaliza con código 0 y cero errores de tipado estricto. Ahora levantamos el entorno de desarrollo con 'pnpm dev'."
    )

    add_bloque(
        "Bloque 3: Flujo del Cliente y Cumplimiento de Reglas de Negocio",
        "2:45 - 5:30",
        [
            "Ir a http://localhost:5173 y hacer clic en 'Ingresar'.",
            "Mostrar la caja de credenciales demo y hacer clic en 'Cliente' (cliente / Cliente1234!).",
            "Iniciar sesión y mostrar el cartel flotante (Toast) de bienvenida.",
            "En el Catálogo: probar buscador por texto, filtro de categorías y paginación en servidor.",
            "Agregar 2 platos al carrito y mostrar cómo se abre el Mini Carrito Deslizable (CartDrawer).",
            "Ir a '/checkout': dejar la dirección con menos de 5 caracteres para ver la validación de TanStack Form (RN-06). Luego ingresar dirección válida, elegir Efectivo y confirmar.",
            "En '/mis-pedidos': mostrar la orden en estado PENDIENTE, desplegar el detalle con precios inmutables (RN-02).",
            "Hacer clic en 'Cancelar Pedido' (RN-05), confirmar en el modal y mostrar cómo el stock se restituye."
        ],
        "Pasamos ahora a la experiencia del cliente. Entramos a la pantalla de login. Fíjense que implementamos TanStack Form con validaciones reactivas y mantuvimos la idea de Prog 3: un recuadro para autocompletar credenciales de prueba con un solo clic.\n\n"
        "Ingresamos como Cliente. El backend establece la cookie HttpOnly y el sistema nos da la bienvenida con una notificación Toast flotante.\n\n"
        "En el catálogo tenemos paginación en servidor mediante skip y limit. Filtramos por categorías y buscamos por texto. Al presionar 'Agregar', se despliega el Mini Carrito Deslizable lateral, permitiendo modificar unidades en vivo respetando el stock disponible.\n\n"
        "Avanzamos al Checkout. Si la dirección tiene menos de 5 caracteres, TanStack Form nos alerta cumpliendo la regla RN-06. Ingresamos una dirección completa, elegimos pago en Efectivo y confirmamos.\n\n"
        "El backend procesa la orden de forma atómica con Unit of Work (regla RN-01), descuenta el stock, vacía el carrito y nos lleva a 'Mis Pedidos'. Aquí vemos la orden en estado PENDIENTE con su foto histórica de precios unitarios inmutables (RN-02). Al presionar 'Cancelar Pedido', se ejecuta la regla RN-05: se valida que la orden esté pendiente, se anula y las cantidades se reintegran al inventario."
    )

    add_bloque(
        "Bloque 4: Flujo del Administrador, Métricas y Máquina de Estados",
        "5:30 - 8:00",
        [
            "Cerrar sesión del cliente y entrar con la cuenta de 'Admin' (admin / Admin1234!).",
            "Pestaña Resumen (/admin/resumen): Mostrar los 6 KPIs (Ingresos entregados, pedidos de hoy, en preparación, pendientes, entregados, cancelados).",
            "Pestaña Pedidos (/admin/pedidos): Filtrar pedidos. Tomar un pedido PENDIENTE y hacer clic en 'Preparar' (pasa a EN_PREPARACION). Luego en 'Marcar Entregado' (RN-03).",
            "Pestaña Categorías (/admin/categorias): Mostrar la tabla con TanStack Table (filtro y paginación en cliente). Crear una categoría con el modal TanStack Form.",
            "Pestaña Productos (/admin/productos): Mostrar ajuste rápido de stock inline (+ y -) y switch de disponibilidad con un clic.",
            "Pestaña Usuarios (/admin/usuarios): Mostrar la lista y señalar que el botón Desactivar para el admin actual está bloqueado ('Protegido')."
        ],
        "Cerramos sesión e ingresamos ahora con el usuario Administrador. El guardián de rutas nos dirige directamente al Panel de Control.\n\n"
        "En 'Resumen' observamos las métricas en tiempo real: facturación total de órdenes entregadas, pedidos del día y estado de cocina.\n\n"
        "En la pestaña 'Pedidos' se aplica la regla RN-03: la máquina de estados. Si una orden está PENDIENTE, el administrador puede pasarla a EN_PREPARACION o cancelarla. Estando en preparación, puede marcarla como ENTREGADO. Una vez entregada o cancelada, la orden entra en estado final y no admite más transiciones.\n\n"
        "En 'Categorías' y 'Productos' implementamos TanStack Table para filtrado y ordenamiento en cliente. Podemos dar de alta platos con TanStack Form y utilizar la edición rápida de stock inline con los botones de más y menos, o pausar la disponibilidad de un plato para que no figure en el catálogo.\n\n"
        "Finalmente, en 'Usuarios', administramos las cuentas. Observen la regla de seguridad: el sistema bloquea la auto-desactivación del administrador logueado, garantizando que el panel nunca quede inaccesible."
    )

    add_bloque(
        "Bloque 5: Puntos Fuertes del Código, Empaquetado y Cierre",
        "8:00 - 9:00",
        [
            "En VS Code: mostrar rápidamente unit_of_work.py (context manager, commit y rollback automático).",
            "Mostrar frontend/src/store/authStore.ts (partialize guardando solo el usuario, sin tokens en localStorage).",
            "En la terminal, ejecutar: powershell -ExecutionPolicy Bypass -File scripts/empaquetar_entrega.ps1.",
            "Mostrar el archivo TPI-Food_Store-Brian_Emanuel_Rios.zip generado de 0.88 MB.",
            "Mostrar el repositorio GitHub con los 11 issues cerrados."
        ],
        "Para finalizar, repasamos dos aspectos fundamentales del código:\n"
        "En 'unit_of_work.py' definimos el context manager con __enter__ y __exit__. Ante cualquier falla o excepción, se ejecuta automáticamente un rollback garantizando atomicidad.\n"
        "En el frontend, 'authStore.ts' persiste únicamente el objeto usuario mediante partialize. La cookie HttpOnly viaja automáticamente con withCredentials, manteniendo el token fuera del alcance de JavaScript.\n\n"
        "Por último, ejecutamos el script 'empaquetar_entrega.ps1', que comprime la solución en un archivo ZIP limpio de 0.88 MB, excluyendo node_modules, entornos virtuales y secretos locales, listo para el campus virtual.\n\n"
        "Todos los cambios fueron sincronizados en GitHub siguiendo Conventional Commits y cerrando los 11 issues correspondientes a los criterios de evaluación.\n\n"
        "Muchas gracias por su atención."
    )

    # Guardar documento
    doc.save(output_path)
    print(f"Documento guardado exitosamente en: {output_path}")

if __name__ == "__main__":
    out = os.path.abspath("Guion_Video_Demostrativo_FoodStore.docx")
    create_guion_docx(out)
